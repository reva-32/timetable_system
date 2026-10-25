"""Word and Excel duty reports with COE/Principal signature blocks.

Two report families are produced:

* Faculty duty reports, attached to the personalised email sent after a
  timetable is published.
* Daily faculty-duty change reports, generated whenever an adjustment is
  approved.

The reports use the authoritative ``assignments`` collection.  Signature
images are optional; when image paths are not configured a clear "DIGITALLY
SIGNED" block with the configured names, titles and signing time is rendered.
"""

from __future__ import annotations

import os
import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from eligibility import normalize_role
from rosters import format_session_label

BASE_DIR = Path(__file__).resolve().parent

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _env_bool(name, default=False):
    raw = os.getenv(name)
    if raw is None or raw == "":
        return default
    return raw.strip().lower() in {"1", "true", "yes", "on"}


def _env_int(name, default):
    try:
        return int(os.getenv(name, str(default)))
    except (TypeError, ValueError):
        return default


def _display_date(value):
    raw = str(value or "").strip()
    for fmt in ("%d-%b-%Y", "%d-%B-%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, fmt).strftime("%d-%b-%Y")
        except ValueError:
            continue
    return raw or "TBD"


def _parse_date(value):
    raw = str(value or "").strip()
    if not raw or raw.upper() in {"TBD", "NONE", "NAN"}:
        return None
    for fmt in ("%d-%b-%Y", "%d-%B-%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(raw.replace("Z", "+00:00")).date()
    except ValueError:
        return None


def _display_datetime(value):
    raw = str(value or "").strip()
    if not raw:
        return "-"
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        return parsed.strftime("%d-%b-%Y, %I:%M %p")
    except ValueError:
        return raw


def _slot_sort(value):
    text = str(value or "")
    match = re.search(r"\d+", text)
    return (int(match.group()) if match else 999, text)


def _safe_component(value, fallback="report"):
    text = re.sub(r"[^A-Za-z0-9._-]+", "_", str(value or "")).strip("._")
    return text or fallback


def _resolve_image_path(name):
    raw = (os.getenv(name) or "").strip()
    if not raw:
        return None
    path = Path(raw)
    if not path.is_absolute():
        path = BASE_DIR / path
    return str(path.resolve()) if path.is_file() else None


def report_identity(signed_at=None):
    """Signature identity shown in every report and email footer."""
    signed_at = signed_at or datetime.now()
    return {
        "institute_name": (
            os.getenv("INSTITUTE_NAME", "").strip()
            or "Pune Institute of Computer Technology (PICT)"
        ),
        "coe_name": (os.getenv("COE_NAME", "").strip() or "Controller of Examinations"),
        "coe_title": (os.getenv("COE_TITLE", "").strip() or "Controller of Examinations (COE)"),
        "principal_name": (os.getenv("PRINCIPAL_NAME", "").strip() or "Principal"),
        "principal_title": (os.getenv("PRINCIPAL_TITLE", "").strip() or "Principal"),
        "coe_image": _resolve_image_path("COE_SIGNATURE_IMAGE"),
        "principal_image": _resolve_image_path("PRINCIPAL_SIGNATURE_IMAGE"),
        "signed_at": signed_at.strftime("%d-%b-%Y, %I:%M %p"),
        "signature_note": "Digitally signed and approved for publication",
    }


def report_output_root():
    raw = (os.getenv("REPORT_OUTPUT_DIR") or "generated_reports").strip()
    path = Path(raw)
    if not path.is_absolute():
        path = BASE_DIR / path
    return path


def _normalized_duty(row):
    return {
        "duty_id": str(row.get("duty_id") or ""),
        "date": _display_date(row.get("date")),
        "date_sort": _parse_date(row.get("date")) or date.max,
        "slot": str(row.get("slot") or ""),
        "session": format_session_label(row.get("session")),
        "course_id": str(row.get("course_id") or "ALL"),
        "year": str(row.get("year") or "ALL"),
        "department": str(row.get("department_required") or row.get("department") or ""),
        "room": str(row.get("room") or row.get("room_assigned") or "TBD"),
        "role": normalize_role(row.get("role") or row.get("role_required") or "Junior"),
    }


def faculty_report_context(timetable, teacher, duty_rows, generated_at=None):
    duties = [_normalized_duty(row) for row in (duty_rows or [])]
    if not teacher or not duties:
        return None
    duties.sort(key=lambda row: (row["date_sort"], _slot_sort(row["slot"]), row["role"]))
    generated_at = generated_at or datetime.now()
    revision = max(1, int(timetable.get("revision") or 1))
    return {
        "kind": "faculty_duty",
        "title": "Faculty Invigilation Duty Report",
        "teacher": {
            "teacher_id": str(teacher.get("teacher_id") or ""),
            "name": teacher.get("name") or str(teacher.get("teacher_id") or "Faculty"),
            "email": str(teacher.get("email") or ""),
            "department": teacher.get("department") or "General",
            "role": normalize_role(teacher.get("role") or "Junior"),
        },
        "timetable": {
            "timetable_id": str(timetable.get("_id") or ""),
            "department": timetable.get("department") or "ALL",
            "version": int(timetable.get("version") or 1),
            "revision": revision,
        },
        "duties": duties,
        "identity": report_identity(generated_at),
        "generated_at": generated_at.strftime("%d-%b-%Y, %I:%M %p"),
        "reporting_minutes": _env_int("JS_REPORTING_MINUTES", 30),
        "squad_reporting_minutes": _env_int("SQUAD_REPORTING_MINUTES", 40),
        "control_room": (os.getenv("CONTROL_ROOM", "Room A1-110").strip() or "Room A1-110"),
    }


def daily_change_context(timetable, report_date, rows, generated_at=None):
    normalized = []
    for row in rows or []:
        normalized.append({
            **row,
            "date_sort": _parse_date(row.get("date")) or date.max,
        })
    if not normalized:
        return None
    normalized.sort(key=lambda row: (_slot_sort(row.get("slot")), str(row.get("previous_teacher_name") or "")))
    generated_at = generated_at or datetime.now()
    revision = max(
        [int(row.get("revision") or 0) for row in normalized] +
        [max(1, int(timetable.get("revision") or 1))]
    )
    return {
        "kind": "daily_change",
        "title": "Daily Faculty Duty Change Report",
        "report_date": _display_date(report_date),
        "timetable": {
            "timetable_id": str(timetable.get("_id") or ""),
            "department": timetable.get("department") or "ALL",
            "version": int(timetable.get("version") or 1),
            "revision": revision,
        },
        "rows": normalized,
        "identity": report_identity(generated_at),
        "generated_at": generated_at.strftime("%d-%b-%Y, %I:%M %p"),
    }


def _find_duty(database, timetable, timetable_id, date_value, slot_value, role=None):
    query = {
        "timetable_id": str(timetable_id),
        "date": _display_date(date_value),
        "slot": str(slot_value),
        "active": True,
    }
    if role:
        query["role"] = normalize_role(role)
    row = database["assignments"].find_one(query)
    if row:
        return row

    wanted_date = _display_date(date_value)
    for candidate in timetable.get("all_duties") or timetable.get("duties") or timetable.get("teacher_duties") or []:
        if (
            _display_date(candidate.get("date")) == wanted_date
            and str(candidate.get("slot")) == str(slot_value)
            and (not role or normalize_role(candidate.get("role") or candidate.get("role_required")) == normalize_role(role))
        ):
            return candidate
    return {}


def load_daily_change_rows(database, timetable, report_date):
    """Build the rows shown in one date's daily change report."""
    timetable_id = str(timetable.get("_id") or "")
    if not timetable_id:
        return []
    target = _parse_date(report_date)
    if not target:
        return []

    teachers = {
        str(doc.get("teacher_id")): doc
        for doc in database["teachers"].find({})
        if doc.get("teacher_id")
    }

    def teacher_name(teacher_id):
        if not teacher_id:
            return "Vacant"
        doc = teachers.get(str(teacher_id)) or {}
        name = doc.get("name") or str(teacher_id)
        return f"{name} ({teacher_id})"

    events = list(database["swap_log"].find({"timetable_id": timetable_id}).sort("created_at", 1))
    rows = []
    for event in events:
        action = str(event.get("action") or "swap").lower()
        if action == "move":
            changes = [
                (event.get("from_date"), event.get("from_slot"), event.get("requester_id"), None, "Moved out"),
                (event.get("to_date"), event.get("to_slot"), None, event.get("requester_id"), "Moved in"),
            ]
        else:
            changes = [
                (event.get("requester_date"), event.get("requester_slot"),
                 event.get("requester_id"), event.get("partner_id"), "Mutual swap"),
                (event.get("partner_date"), event.get("partner_slot"),
                 event.get("partner_id"), event.get("requester_id"), "Mutual swap"),
            ]

        for change_date, change_slot, previous_id, new_id, label in changes:
            if _parse_date(change_date) != target:
                continue
            duty = _find_duty(
                database, timetable, timetable_id, change_date, change_slot,
                event.get("role"),
            )
            duty = _normalized_duty(duty)
            actor = teachers.get(str(event.get("actor_id") or "")) or {}
            rows.append({
                "date": _display_date(change_date),
                "slot": str(change_slot or ""),
                "session": duty["session"],
                "course_id": duty["course_id"],
                "year": duty["year"],
                "room": duty["room"],
                "role": duty["role"],
                "previous_teacher_id": str(previous_id or ""),
                "previous_teacher_name": teacher_name(previous_id),
                "new_teacher_id": str(new_id or ""),
                "new_teacher_name": teacher_name(new_id),
                "change": label,
                "reason": event.get("reason") or "-",
                "actor": actor.get("name") or event.get("actor_id") or "-",
                "approved_at": _display_datetime(event.get("created_at")),
                "revision": int(event.get("revision") or timetable.get("revision") or 1),
            })
    return rows


def save_report_buffers(kind, timetable_id, stem, buffers):
    """Persist report bytes and return file metadata."""
    folder = report_output_root() / _safe_component(kind) / _safe_component(timetable_id)
    folder.mkdir(parents=True, exist_ok=True)
    saved = {}
    for extension, payload in buffers.items():
        filename = f"{_safe_component(stem)}_{_safe_component(timetable_id)[-6:]}.{extension}"
        path = folder / filename
        path.write_bytes(payload)
        saved[extension] = {"filename": filename, "path": str(path)}
    return saved


def _docx_signature_block(document, identity, heading_size=8):
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    table = document.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    parties = [
        (identity["coe_name"], identity["coe_title"], identity["coe_image"]),
        (identity["principal_name"], identity["principal_title"], identity["principal_image"]),
    ]
    for index, (name, title, image_path) in enumerate(parties):
        cell = table.cell(0, index)
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "EAF2FF")
        tc_pr.append(shading)

        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("DIGITALLY SIGNED")
        run.bold = True
        run.font.size = Pt(heading_size)
        run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)

        image_cell = table.cell(1, index)
        image_paragraph = image_cell.paragraphs[0]
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if image_path:
            image_paragraph.add_run().add_picture(image_path, width=Inches(1.45))
        else:
            fallback = image_paragraph.add_run("Digitally verified")
            fallback.italic = True
            fallback.font.size = Pt(8)
            fallback.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        name_paragraph = table.cell(2, index).paragraphs[0]
        name_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name_run = name_paragraph.add_run(name)
        name_run.bold = True
        name_run.font.size = Pt(10)

        title_paragraph = table.cell(3, index).paragraphs[0]
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_paragraph.add_run(title)
        title_run.font.size = Pt(8.5)

        date_paragraph = table.cell(4, index).paragraphs[0]
        date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_run = date_paragraph.add_run(identity["signed_at"])
        date_run.font.size = Pt(8)
        date_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)


def _docx_cell_text(cell, value, bold=False, size=8.5, color=None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(str(value))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)


def _docx_shade(cell, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _docx_header(document, context):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    identity = context["identity"]
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(identity["institute_name"])
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run("EXAMINATION SECTION  •  INVIGILATION DUTY REPORT")
    sub_run.bold = True
    sub_run.font.size = Pt(9)
    sub_run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(context["title"])
    title_run.bold = True
    title_run.font.size = Pt(13)
    title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)


def _add_docx_footer(document):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(
        "PICT Examination Section  •  System-generated report  •  "
        "Digitally approved by COE and Principal"
    )
    run.font.size = Pt(7.5)
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)


def render_faculty_docx(context):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(9.5)
    document.core_properties.title = context["title"]
    document.core_properties.author = context["identity"]["institute_name"]
    document.core_properties.subject = "Faculty invigilation duties"

    _docx_header(document, context)

    teacher = context["teacher"]
    timetable = context["timetable"]
    meta = document.add_table(rows=3, cols=2)
    meta.style = "Table Grid"
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_values = [
        (f"Faculty: {teacher['name']} ({teacher['teacher_id']})", f"Department: {teacher['department']}"),
        (f"Role: {teacher['role']}", f"Department roster: {timetable['department']}"),
        (
            f"Timetable version: {timetable['version']}    |    Revision: {timetable['revision']}",
            f"Date of issue: {context['generated_at']}",
        ),
    ]
    for row_index, values in enumerate(meta_values):
        for col_index, value in enumerate(values):
            cell = meta.cell(row_index, col_index)
            _docx_cell_text(cell, value, size=8.5)
            _docx_shade(cell, "F1F5F9")

    document.add_paragraph()
    intro = document.add_paragraph()
    intro.add_run(
        "You are assigned the following examination invigilation duties. "
        "Please report to the Examination Control Room on time."
    ).font.size = Pt(9)

    table = document.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["#", "Date", "Slot", "Session", "Course", "Year", "Room", "Role"]
    widths = [0.35, 1.0, 0.55, 1.25, 0.75, 0.65, 0.75, 0.75]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _docx_cell_text(cell, header, bold=True, size=8.5, color=(0xFF, 0xFF, 0xFF))
        _docx_shade(cell, "1E3A5F")
        cell.width = Inches(widths[index])

    for index, duty in enumerate(context["duties"], start=1):
        row = table.add_row()
        values = [
            index, duty["date"], duty["slot"], duty["session"], duty["course_id"],
            duty["year"], duty["room"], duty["role"],
        ]
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            _docx_cell_text(cell, value, size=8.5)
            cell.width = Inches(widths[col_index])
            if index % 2 == 0:
                _docx_shade(cell, "F8FAFC")

    notes = document.add_paragraph()
    notes.paragraph_format.space_before = Pt(8)
    note_run = notes.add_run(
        f"Reporting instructions: report at least {context['reporting_minutes']} minutes before the session "
        f"for Junior Supervisor duties and {context['squad_reporting_minutes']} minutes before the session "
        f"for Squad duties, to the Examination Control Room ({context['control_room']})."
    )
    note_run.font.size = Pt(8.5)
    note_run.italic = True
    note_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    document.add_paragraph()
    signature_heading = document.add_paragraph()
    signature_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    signature_run = signature_heading.add_run(context["identity"]["signature_note"])
    signature_run.bold = True
    signature_run.font.size = Pt(8.5)
    signature_run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
    _docx_signature_block(document, context["identity"])
    _add_docx_footer(document)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_faculty_xlsx(context):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    identity = context["identity"]
    teacher = context["teacher"]
    timetable = context["timetable"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Duty Report"
    sheet.sheet_view.showGridLines = False

    last_column = 8
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    accent_fill = PatternFill("solid", fgColor="EAF2FF")
    light_fill = PatternFill("solid", fgColor="F8FAFC")

    def merged_row(row_number, value, size=9, bold=False, fill=None, color="0F172A", height=20):
        sheet.merge_cells(
            start_row=row_number, start_column=1,
            end_row=row_number, end_column=last_column,
        )
        cell = sheet.cell(row=row_number, column=1, value=value)
        cell.font = Font(size=size, bold=bold, color=color)
        cell.alignment = center
        if fill:
            cell.fill = fill
        for column in range(1, last_column + 1):
            sheet.cell(row=row_number, column=column).border = border
        sheet.row_dimensions[row_number].height = height

    merged_row(1, identity["institute_name"], size=16, bold=True, color="1E3A5F", height=26)
    merged_row(2, "EXAMINATION SECTION  •  INVIGILATION DUTY REPORT", size=9, bold=True, color="0F766E")
    merged_row(3, context["title"], size=13, bold=True, color="1E3A5F", height=24)
    merged_row(
        4,
        f"Faculty: {teacher['name']} ({teacher['teacher_id']})    |    "
        f"Department: {teacher['department']}    |    Role: {teacher['role']}",
        size=9,
    )
    merged_row(
        5,
        f"Timetable: {timetable['department']}    |    Version: {timetable['version']}    |    "
        f"Revision: {timetable['revision']}    |    Issued: {context['generated_at']}",
        size=8.5,
    )
    merged_row(
        6,
        "Please report to the Examination Control Room on time for every duty listed below.",
        size=8.5,
        fill=accent_fill,
    )

    header_row = 8
    headers = ["#", "Date", "Slot", "Session", "Course", "Year", "Room", "Role"]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, size=9, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    sheet.row_dimensions[header_row].height = 24

    row_number = header_row + 1
    for index, duty in enumerate(context["duties"], start=1):
        values = [
            index, duty["date"], duty["slot"], duty["session"], duty["course_id"],
            duty["year"], duty["room"], duty["role"],
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            cell.font = Font(size=9)
            cell.alignment = center if column != 4 else left
            cell.border = border
            if index % 2 == 0:
                cell.fill = light_fill
        row_number += 1

    row_number += 1
    sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=last_column)
    note = sheet.cell(
        row=row_number,
        column=1,
        value=(
            f"Reporting: at least {context['reporting_minutes']} minutes before the session for Junior Supervisor "
            f"duties and {context['squad_reporting_minutes']} minutes before the session for Squad duties. "
            f"Control Room: {context['control_room']}."
        ),
    )
    note.font = Font(size=8.5, italic=True, color="475569")
    note.alignment = left
    note.fill = PatternFill("solid", fgColor="F8FAFC")

    row_number += 2
    _excel_signature_block(sheet, row_number, identity, first_column=1, last_column=4)
    _excel_signature_block(sheet, row_number, identity, first_column=5, last_column=8)

    widths = [5, 14, 8, 25, 13, 10, 12, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    sheet.print_title_rows = f"1:{header_row}"
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.center.text = "PICT Examination Section - Page &P of &N"
    sheet.oddFooter.center.size = 8

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _excel_signature_block(sheet, start_row, identity, first_column, last_column):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="93C5FD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    heading_fill = PatternFill("solid", fgColor="EAF2FF")
    first_col_letter = sheet.cell(row=start_row, column=first_column).column_letter
    last_col_letter = sheet.cell(row=start_row, column=last_column).column_letter

    sheet.merge_cells(
        start_row=start_row, start_column=first_column,
        end_row=start_row, end_column=last_column,
    )
    heading = sheet.cell(row=start_row, column=first_column, value="DIGITALLY SIGNED")
    heading.font = Font(size=8, bold=True, color="1D4ED8")
    heading.fill = heading_fill
    heading.alignment = center

    sheet.merge_cells(
        start_row=start_row + 1, start_column=first_column,
        end_row=start_row + 2, end_column=last_column,
    )
    if identity["coe_image"] and first_column <= 4:
        image_path = identity["coe_image"]
    elif identity["principal_image"] and first_column > 4:
        image_path = identity["principal_image"]
    else:
        image_path = None

    if image_path:
        from openpyxl.drawing.image import Image as ExcelImage
        image = ExcelImage(image_path)
        image.width = 135
        image.height = 45
        sheet.add_image(image, f"{first_col_letter}{start_row + 1}")
        signature_cell = sheet.cell(row=start_row + 1, column=first_column)
        signature_cell.value = ""
    else:
        signature_cell = sheet.cell(row=start_row + 1, column=first_column, value="Digitally verified")
        signature_cell.font = Font(size=8, italic=True, color="64748B")
        signature_cell.alignment = center

    party_name = identity["coe_name"] if first_column <= 4 else identity["principal_name"]
    party_title = identity["coe_title"] if first_column <= 4 else identity["principal_title"]
    rows = [
        (start_row + 3, party_name, 10, True, "0F172A"),
        (start_row + 4, party_title, 8.5, False, "334155"),
        (start_row + 5, identity["signed_at"], 8, False, "64748B"),
    ]
    for row_number, value, size, bold, color in rows:
        sheet.merge_cells(
            start_row=row_number, start_column=first_column,
            end_row=row_number, end_column=last_column,
        )
        cell = sheet.cell(row=row_number, column=first_column, value=value)
        cell.font = Font(size=size, bold=bold, color=color)
        cell.alignment = center

    for row_number in range(start_row, start_row + 6):
        for column in range(first_column, last_column + 1):
            sheet.cell(row=row_number, column=column).border = border
    for offset in (1, 2):
        sheet.row_dimensions[start_row + offset].height = 26


def render_daily_change_docx(context):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(8.5)
    document.core_properties.title = context["title"]
    document.core_properties.author = context["identity"]["institute_name"]

    _docx_header(document, context)
    report_date_paragraph = document.add_paragraph()
    report_date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    report_date_run = report_date_paragraph.add_run(f"Examination Date: {context['report_date']}")
    report_date_run.bold = True
    report_date_run.font.size = Pt(10)
    report_date_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

    meta = document.add_table(rows=1, cols=3)
    meta.style = "Table Grid"
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = [
        f"Department: {context['timetable']['department']}",
        f"Timetable version: {context['timetable']['version']}    |    Revision: {context['timetable']['revision']}",
        f"Generated: {context['generated_at']}",
    ]
    for index, value in enumerate(values):
        cell = meta.cell(0, index)
        _docx_cell_text(cell, value, size=8.5)
        _docx_shade(cell, "F1F5F9")

    document.add_paragraph()
    table = document.add_table(rows=1, cols=11)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = [
        "#", "Slot", "Session", "Course", "Year", "Room",
        "Previous Faculty", "New Faculty", "Change", "Reason", "Approved At",
    ]
    widths = [0.3, 0.5, 1.15, 0.65, 0.55, 0.65, 1.35, 1.35, 0.75, 1.55, 1.1]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _docx_cell_text(cell, header, bold=True, size=7.5, color=(0xFF, 0xFF, 0xFF))
        _docx_shade(cell, "1E3A5F")
        cell.width = Inches(widths[index])

    for index, row_data in enumerate(context["rows"], start=1):
        row = table.add_row()
        values = [
            index,
            row_data.get("slot", ""),
            row_data.get("session", ""),
            row_data.get("course_id", ""),
            row_data.get("year", ""),
            row_data.get("room", ""),
            row_data.get("previous_teacher_name", ""),
            row_data.get("new_teacher_name", ""),
            row_data.get("change", ""),
            row_data.get("reason", ""),
            row_data.get("approved_at", ""),
        ]
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            _docx_cell_text(cell, value, size=7.5)
            cell.width = Inches(widths[col_index])
            if index % 2 == 0:
                _docx_shade(cell, "F8FAFC")

    summary = document.add_paragraph()
    summary.paragraph_format.space_before = Pt(8)
    summary_run = summary.add_run(f"Total duty changes on this date: {len(context['rows'])}")
    summary_run.bold = True
    summary_run.font.size = Pt(9)
    summary_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

    note = document.add_paragraph(
        "This report is generated from the published assignment and swap log. "
        "It should be retained with the signed examination records."
    )
    note_run = note.runs[0]
    note_run.italic = True
    note_run.font.size = Pt(8)
    note_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    document.add_paragraph()
    signature_heading = document.add_paragraph()
    signature_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    signature_run = signature_heading.add_run(context["identity"]["signature_note"])
    signature_run.bold = True
    signature_run.font.size = Pt(8.5)
    signature_run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
    _docx_signature_block(document, context["identity"])
    _add_docx_footer(document)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_daily_change_xlsx(context):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    identity = context["identity"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Daily Changes"
    sheet.sheet_view.showGridLines = False

    last_column = 11
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    accent_fill = PatternFill("solid", fgColor="EAF2FF")
    light_fill = PatternFill("solid", fgColor="F8FAFC")

    def merged_row(row_number, value, size=9, bold=False, fill=None, color="0F172A", height=20):
        sheet.merge_cells(
            start_row=row_number, start_column=1,
            end_row=row_number, end_column=last_column,
        )
        cell = sheet.cell(row=row_number, column=1, value=value)
        cell.font = Font(size=size, bold=bold, color=color)
        cell.alignment = center
        if fill:
            cell.fill = fill
        for column in range(1, last_column + 1):
            sheet.cell(row=row_number, column=column).border = border
        sheet.row_dimensions[row_number].height = height

    merged_row(1, identity["institute_name"], size=16, bold=True, color="1E3A5F", height=26)
    merged_row(2, "EXAMINATION SECTION  •  DAILY FACULTY DUTY CHANGE REPORT", size=9, bold=True, color="0F766E")
    merged_row(3, f"{context['title']} — {context['report_date']}", size=13, bold=True, color="1E3A5F", height=24)
    merged_row(
        4,
        f"Department: {context['timetable']['department']}    |    "
        f"Version: {context['timetable']['version']}    |    Revision: {context['timetable']['revision']}    |    "
        f"Generated: {context['generated_at']}",
        size=8.5,
    )
    merged_row(
        5,
        f"Total duty changes on this date: {len(context['rows'])}",
        size=9,
        bold=True,
        fill=accent_fill,
    )

    header_row = 7
    headers = [
        "#", "Slot", "Session", "Course", "Year", "Room",
        "Previous Faculty", "New Faculty", "Change", "Reason", "Approved At",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, size=8.5, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    sheet.row_dimensions[header_row].height = 26

    row_number = header_row + 1
    for index, row_data in enumerate(context["rows"], start=1):
        values = [
            index,
            row_data.get("slot", ""),
            row_data.get("session", ""),
            row_data.get("course_id", ""),
            row_data.get("year", ""),
            row_data.get("room", ""),
            row_data.get("previous_teacher_name", ""),
            row_data.get("new_teacher_name", ""),
            row_data.get("change", ""),
            row_data.get("reason", ""),
            row_data.get("approved_at", ""),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            cell.font = Font(size=8.5)
            cell.alignment = left if column in {3, 7, 8, 10, 11} else center
            cell.border = border
            if index % 2 == 0:
                cell.fill = light_fill
        row_number += 1

    row_number += 2
    _excel_signature_block(sheet, row_number, identity, first_column=1, last_column=5)
    _excel_signature_block(sheet, row_number, identity, first_column=6, last_column=11)

    widths = [5, 7, 22, 12, 9, 10, 22, 22, 14, 28, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    sheet.print_title_rows = f"1:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.center.text = "PICT Examination Section - Page &P of &N"
    sheet.oddFooter.center.size = 8

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def faculty_report_filename(context, extension):
    teacher = context["teacher"]
    timetable = context["timetable"]
    return (
        f"Faculty_Duty_{_safe_component(teacher['teacher_id'])}_"
        f"{_safe_component(timetable['department'])}_Rev{timetable['revision']}.{extension}"
    )


def daily_change_filename(context, extension):
    timetable = context["timetable"]
    iso_date = _parse_date(context["report_date"])
    date_part = iso_date.strftime("%Y-%m-%d") if iso_date else _safe_component(context["report_date"])
    return f"Daily_Change_{date_part}_{_safe_component(timetable['department'])}_Rev{timetable['revision']}.{extension}"


def generate_faculty_reports(database, timetable, teacher, persist=True, generated_at=None):
    timetable_id = str(timetable.get("_id") or "")
    teacher_id = str(teacher.get("teacher_id") or "")
    if not timetable_id or not teacher_id:
        return None
    duties = list(database["assignments"].find({
        "timetable_id": timetable_id,
        "teacher_id": teacher_id,
        "active": True,
    }))
    context = faculty_report_context(timetable, teacher, duties, generated_at=generated_at)
    if not context:
        return None
    docx_bytes = render_faculty_docx(context)
    xlsx_bytes = render_faculty_xlsx(context)
    stem = (
        f"Faculty_Duty_{_safe_component(teacher_id)}_"
        f"{_safe_component(context['timetable']['department'])}_Rev{context['timetable']['revision']}"
    )
    files = {}
    if persist:
        files = save_report_buffers(
            "faculty_duties", timetable_id, stem,
            {"docx": docx_bytes, "xlsx": xlsx_bytes},
        )
    return {
        "context": context,
        "buffers": {"docx": docx_bytes, "xlsx": xlsx_bytes},
        "filenames": {
            "docx": faculty_report_filename(context, "docx"),
            "xlsx": faculty_report_filename(context, "xlsx"),
        },
        "files": files,
    }


def generate_daily_change_reports(database, timetable, report_dates, persist=True,
                                  generated_at=None, generated_by=None):
    timetable_id = str(timetable.get("_id") or "")
    if not timetable_id:
        return []
    unique_dates = []
    seen = set()
    for value in report_dates or []:
        parsed = _parse_date(value)
        key = parsed.isoformat() if parsed else str(value)
        if key not in seen:
            seen.add(key)
            unique_dates.append(parsed or value)

    generated = []
    for report_date in unique_dates:
        rows = load_daily_change_rows(database, timetable, report_date)
        context = daily_change_context(timetable, report_date, rows, generated_at=generated_at)
        if not context:
            continue
        docx_bytes = render_daily_change_docx(context)
        xlsx_bytes = render_daily_change_xlsx(context)
        parsed = _parse_date(context["report_date"])
        iso_date = parsed.isoformat() if parsed else _safe_component(context["report_date"])
        stem = (
            f"Daily_Change_{iso_date}_{_safe_component(context['timetable']['department'])}_"
            f"Rev{context['timetable']['revision']}"
        )
        files = {}
        if persist:
            files = save_report_buffers(
                "daily_changes", timetable_id, stem,
                {"docx": docx_bytes, "xlsx": xlsx_bytes},
            )
            try:
                database["change_reports"].update_one(
                    {
                        "timetable_id": timetable_id,
                        "report_date": iso_date,
                        "revision": context["timetable"]["revision"],
                    },
                    {"$set": {
                        "timetable_id": timetable_id,
                        "report_date": iso_date,
                        "report_date_display": context["report_date"],
                        "revision": context["timetable"]["revision"],
                        "row_count": len(context["rows"]),
                        "generated_at": context["generated_at"],
                        "generated_by": str(generated_by) if generated_by else None,
                        "files": files,
                    }},
                    upsert=True,
                )
            except Exception:
                pass
        generated.append({
            "context": context,
            "buffers": {"docx": docx_bytes, "xlsx": xlsx_bytes},
            "filenames": {
                "docx": daily_change_filename(context, "docx"),
                "xlsx": daily_change_filename(context, "xlsx"),
            },
            "files": files,
        })
    return generated
