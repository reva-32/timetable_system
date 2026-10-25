"""College-format invigilation rosters (PDF and Excel).

The college publishes two matrices per examination: Junior Supervisors and
Internal Squad.  Both are generated from the authoritative ``assignments``
collection so a revision always reflects approved swaps.
"""

from __future__ import annotations

import os
import re
from datetime import date, datetime
from io import BytesIO

from eligibility import normalize_role, session_code, session_start_hour

ROSTER_TYPES = {
    "junior": ("Junior", "Junior Supervisor Duty Schedule"),
    "js": ("Junior", "Junior Supervisor Duty Schedule"),
    "junior-supervisor": ("Junior", "Junior Supervisor Duty Schedule"),
    "junior_supervisor": ("Junior", "Junior Supervisor Duty Schedule"),
    "squad": ("Squad", "Internal Squad Supervisor Duty Schedule"),
    "internal-squad": ("Squad", "Internal Squad Supervisor Duty Schedule"),
    "internal_squad": ("Squad", "Internal Squad Supervisor Duty Schedule"),
    "is": ("Squad", "Internal Squad Supervisor Duty Schedule"),
}


def normalize_roster_type(value):
    key = str(value or "").strip().lower().replace(" ", "-")
    if key not in ROSTER_TYPES:
        raise ValueError("Unknown roster type. Use 'junior' or 'squad'.")
    return key


def roster_role_and_title(roster_type):
    return ROSTER_TYPES[normalize_roster_type(roster_type)]


def _env_int(name, default):
    try:
        return max(0, int(os.getenv(name, str(default))))
    except (TypeError, ValueError):
        return default


def reporting_text(roster_type):
    role, _ = roster_role_and_title(roster_type)
    minutes = _env_int(
        "SQUAD_REPORTING_MINUTES" if role == "Squad" else "JS_REPORTING_MINUTES",
        40 if role == "Squad" else 30,
    )
    room = os.getenv("CONTROL_ROOM", "Room A1-110").strip() or "Room A1-110"
    return f"Reporting time: {minutes} minutes before the session in the Examination Control Room ({room})."


def _parse_date(value):
    text = str(value or "").strip()
    if not text or text.upper() in {"TBD", "NONE", "NAN"}:
        return None
    for fmt in ("%d-%b-%Y", "%d-%B-%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        parsed = datetime.fromisoformat(text)
        return parsed.date()
    except ValueError:
        return None


def _display_date(value):
    parsed = _parse_date(value)
    return parsed.strftime("%d-%b-%Y") if parsed else str(value)


def _month_year(dates):
    months = sorted({(d.year, d.month) for d in dates if d})
    if not months:
        return date.today().strftime("%B %Y")
    labels = [f"{date(year, month, 1):%B %Y}" for year, month in months]
    if len(labels) == 1:
        return labels[0]
    return f"{labels[0]} - {labels[-1]}"


def _clock(value, meridiem=None):
    """Format HH:MM as a 12-hour clock label, applying AM/PM when supplied."""
    try:
        hour, minute = [int(part) for part in value.split(":")[:2]]
    except (TypeError, ValueError):
        return None
    if meridiem:
        meridiem = meridiem.lower()
        if meridiem == "pm" and hour < 12:
            hour += 12
        elif meridiem == "am" and hour == 12:
            hour = 0
    suffix = "am" if hour < 12 else "pm"
    display_hour = hour % 12 or 12
    return f"{display_hour}:{minute:02d} {suffix}"


def format_session_label(session):
    """Turn a time range or name into the college roster session line."""
    text = str(session or "").strip()
    if not text:
        return "Session: TBD"

    lower = text.lower()
    label = None
    for key in ("morning", "afternoon", "evening"):
        if lower.startswith(key):
            label = key.capitalize()
            break

    match = re.search(
        r"(\d{1,2}:\d{2})\s*(am|pm)?\s*(?:-|–|—|to)\s*(\d{1,2}:\d{2})\s*(am|pm)?",
        text,
        flags=re.IGNORECASE,
    )
    if match:
        start, start_meridiem, end, end_meridiem = match.groups()
        start_label = _clock(start, start_meridiem)
        end_label = _clock(end, end_meridiem)
        if start_label and end_label:
            if label:
                return f"{label}: {start_label} to {end_label}"
            hour = session_start_hour(text)
            if hour is None:
                hour = 10
            period = "Morning" if hour < 12 else ("Afternoon" if hour < 17 else "Evening")
            return f"{period}: {start_label} to {end_label}"
    if label:
        return label
    return text


def _session_codes(session):
    return session_code(session)


def unique_session_lines(rows):
    seen = []
    for row in rows or []:
        label = format_session_label(row.get("session"))
        if label and label not in seen:
            seen.append(label)
    return seen


def _teacher_serials(database, timetable_id, department, role, teachers):
    """Stable Sr. No. mapping, preserved across roster revisions."""
    eligible = [
        teacher for teacher in teachers or []
        if normalize_role(teacher.get("role")) == role
        and not teacher.get("is_admin")
        and teacher.get("teacher_id")
    ]
    eligible.sort(key=lambda item: (str(item.get("teacher_id")), str(item.get("name") or "")))
    ids = [str(teacher.get("teacher_id")) for teacher in eligible]

    serials = {}
    snapshot = None
    if database is not None and timetable_id:
        try:
            snapshot = database["roster_snapshots"].find_one({
                "timetable_id": str(timetable_id),
                "roster_type": role.lower(),
            })
        except Exception:
            snapshot = None
    if snapshot and isinstance(snapshot.get("serials"), dict):
        serials = {str(key): int(value) for key, value in snapshot["serials"].items() if str(value).isdigit()}

    next_serial = max(serials.values(), default=0) + 1
    for teacher_id in ids:
        if teacher_id not in serials:
            serials[teacher_id] = next_serial
            next_serial += 1

    if database is not None and timetable_id:
        try:
            database["roster_snapshots"].update_one(
                {"timetable_id": str(timetable_id), "roster_type": role.lower()},
                {"$set": {
                    "timetable_id": str(timetable_id),
                    "roster_type": role.lower(),
                    "department": department,
                    "serials": serials,
                    "updated_at": datetime.utcnow().isoformat() + "Z",
                }},
                upsert=True,
            )
        except Exception:
            pass

    return {teacher_id: serials[teacher_id] for teacher_id in ids}


def build_roster_context(timetable, assignments, teachers, roster_type,
                         database=None, revision=None, issue_date=None):
    """Build the matrix data shared by the PDF and Excel renderers."""
    role, title = roster_role_and_title(roster_type)
    timetable = timetable or {}
    timetable_id = str(timetable.get("_id") or "")
    department = timetable.get("department") or "ALL"

    active_rows = [
        row for row in (assignments or [])
        if normalize_role(row.get("role")) == role and row.get("teacher_id")
    ]

    date_values = []
    for row in (timetable.get("timetable") or []):
        parsed = _parse_date(row.get("date"))
        if parsed and parsed not in date_values:
            date_values.append(parsed)
    for row in active_rows:
        parsed = _parse_date(row.get("date"))
        if parsed and parsed not in date_values:
            date_values.append(parsed)
    date_values.sort()
    display_dates = [value.strftime("%d-%b-%Y") for value in date_values]

    # Session codes shown in the second header row.  Use every exam on the
    # date, not just this role's duties, so the matrix explains the day.
    codes_by_date = {}
    for row in (timetable.get("timetable") or []):
        parsed = _parse_date(row.get("date"))
        if not parsed:
            continue
        code = _session_codes(row.get("session"))
        codes_by_date.setdefault(parsed.strftime("%d-%b-%Y"), [])
        if code not in codes_by_date[parsed.strftime("%d-%b-%Y")]:
            codes_by_date[parsed.strftime("%d-%b-%Y")].append(code)
    for key in codes_by_date:
        codes_by_date[key].sort()

    serials = _teacher_serials(database, timetable_id, department, role, teachers)
    by_teacher = {}
    for row in active_rows:
        tid = str(row.get("teacher_id"))
        by_teacher.setdefault(tid, []).append(row)

    roster_rows = []
    for teacher in teachers or []:
        if normalize_role(teacher.get("role")) != role or teacher.get("is_admin"):
            continue
        tid = str(teacher.get("teacher_id") or "")
        if not tid:
            continue
        cells = {}
        for row in by_teacher.get(tid, []):
            key = _display_date(row.get("date"))
            code = _session_codes(row.get("session"))
            cells.setdefault(key, [])
            if code not in cells[key]:
                cells[key].append(code)
        roster_rows.append({
            "teacher_id": tid,
            "name": teacher.get("name") or f"Prof. {tid}",
            "serial": serials.get(tid),
            "cells": {key: "/".join(sorted(values)) for key, values in cells.items()},
        })
    roster_rows.sort(key=lambda item: (item["serial"] is None, item["serial"] or 10**9))

    if revision is None:
        revision = int(timetable.get("revision") or 1)
    revision = max(1, int(revision))

    return {
        "role": role,
        "title": title,
        "month_year": _month_year(date_values),
        "dates": display_dates,
        "codes_by_date": codes_by_date,
        "rows": roster_rows,
        "reporting_text": reporting_text(roster_type),
        "session_lines": unique_session_lines(timetable.get("timetable") or []),
        "revision": revision,
        "issue_date": (issue_date or date.today()).strftime("%d-%b-%Y"),
        "department": department,
        "timetable_id": timetable_id,
        "reporting_minutes": _env_int(
            "SQUAD_REPORTING_MINUTES" if role == "Squad" else "JS_REPORTING_MINUTES",
            40 if role == "Squad" else 30,
        ),
        "control_room": os.getenv("CONTROL_ROOM", "Room A1-110").strip() or "Room A1-110",
    }


def _safe_text(value):
    """ReportLab's built-in fonts are Latin-1; keep the roster portable."""
    return (
        str(value)
        .replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u2018", "'")
        .replace("\u2019", "'")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
    )


def render_pdf(context):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdf_canvas
    from reportlab.platypus import (
        KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    class NumberedCanvas(pdf_canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_pages = []

        def showPage(self):
            self._saved_pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._saved_pages)
            for state in self._saved_pages:
                self.__dict__.update(state)
                self.setFont("Helvetica", 7)
                self.setFillColor(colors.HexColor("#475569"))
                self.drawRightString(
                    self._pagesize[0] - 12 * mm, 8 * mm,
                    f"Page {self._pageNumber} of {total}",
                )
                self.drawString(12 * mm, 8 * mm, "ExamSched - Examination Control Room")
                super().showPage()
            super().save()

    buffer = BytesIO()
    page_width, _ = landscape(A4)
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=16 * mm,
        title=_safe_text(context["title"]),
        author="ExamSched",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "RosterTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=13, leading=16, alignment=TA_CENTER, spaceAfter=2,
    )
    subtitle_style = ParagraphStyle(
        "RosterSubtitle", parent=styles["Normal"], fontName="Helvetica",
        fontSize=8.5, leading=11, alignment=TA_CENTER,
    )
    cell_style = ParagraphStyle(
        "RosterCell", parent=styles["Normal"], fontName="Helvetica",
        fontSize=7.2, leading=8.4,
    )
    head_style = ParagraphStyle(
        "RosterHead", parent=cell_style, fontName="Helvetica-Bold",
        alignment=TA_CENTER,
    )
    name_head_style = ParagraphStyle(
        "RosterNameHead", parent=head_style, alignment=TA_LEFT,
    )

    story = []
    story.append(Paragraph("PUNE INSTITUTE OF COMPUTER TECHNOLOGY", subtitle_style))
    story.append(Paragraph(
        _safe_text(f"SPPU BE In Semester Examination {context['month_year']}"),
        subtitle_style,
    ))
    story.append(Paragraph(_safe_text(context["title"]), title_style))
    story.append(Paragraph(_safe_text(context["reporting_text"]), subtitle_style))
    session_text = "   |   ".join(context["session_lines"]) or "Session: TBD"
    story.append(Paragraph(_safe_text(f"Sessions: {session_text}"), subtitle_style))

    revision_label = f"Revision {context['revision']}" if context["revision"] > 1 else "Original issue"
    meta = (
        f"Department: {context['department']}    |    "
        f"{revision_label}    |    Date of issue: {context['issue_date']}"
    )
    story.append(Paragraph(_safe_text(meta), subtitle_style))
    story.append(Spacer(1, 4 * mm))

    dates = context["dates"]
    header = [
        Paragraph("Sr. No.", head_style),
        Paragraph("Name of the Supervisor", name_head_style),
    ] + [Paragraph(_safe_text(value), head_style) for value in dates]
    code_row = [Paragraph("", head_style), Paragraph("Session code", head_style)]
    code_row += [
        Paragraph(_safe_text("/".join(context["codes_by_date"].get(value, []))), head_style)
        for value in dates
    ]

    body = [header, code_row]
    for row in context["rows"]:
        body.append([
            Paragraph(str(row["serial"] or ""), head_style),
            Paragraph(_safe_text(row["name"]), cell_style),
        ] + [
            Paragraph(_safe_text(row["cells"].get(day, "")), head_style)
            for day in dates
        ])

    if len(dates) <= 8:
        date_width = 20 * mm
    elif len(dates) <= 14:
        date_width = 15 * mm
    elif len(dates) <= 22:
        date_width = 11 * mm
    else:
        date_width = max(7.5 * mm, (page_width - 85 * mm) / max(1, len(dates)))
    col_widths = [11 * mm, 48 * mm] + [date_width] * len(dates)

    table = Table(body, colWidths=col_widths, repeatRows=2, hAlign="CENTER")
    box = colors.HexColor("#334155")
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, box),
        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#e2e8f0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("ALIGN", (1, 2), (1, -1), "LEFT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("ROWBACKGROUNDS", (0, 2), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
    ]))
    story.append(table)

    story.append(Spacer(1, 10 * mm))
    signature = Table(
        [[
            Paragraph("____________________________<br/><br/>Senior Supervisor", subtitle_style),
            Paragraph("____________________________<br/><br/>Controller of Examinations (COE)", subtitle_style),
            Paragraph("____________________________<br/><br/>Principal", subtitle_style),
        ]],
        colWidths=[(page_width - 20 * mm) / 3.0] * 3,
    )
    signature.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(KeepTogether(signature))

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer


def render_excel(context):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Junior" if context["role"] == "Junior" else "Internal Squad"

    dates = context["dates"]
    last_column = 2 + len(dates)
    thin = Side(style="thin", color="334155")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="E2E8F0")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(last_column, 2))
    cell = sheet.cell(row=1, column=1, value="SPPU BE In Semester Examination " + context["month_year"])
    cell.font = Font(bold=True, size=14)
    cell.alignment = center

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(last_column, 2))
    cell = sheet.cell(row=2, column=1, value=context["title"])
    cell.font = Font(bold=True, size=12)
    cell.alignment = center

    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max(last_column, 2))
    cell = sheet.cell(row=3, column=1, value=context["reporting_text"])
    cell.font = Font(size=9)
    cell.alignment = left

    session_text = "   |   ".join(context["session_lines"]) or "Session: TBD"
    sheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=max(last_column, 2))
    cell = sheet.cell(row=4, column=1, value="Sessions: " + session_text)
    cell.font = Font(size=9)
    cell.alignment = left

    revision_label = f"Revision {context['revision']}" if context["revision"] > 1 else "Original issue"
    sheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=max(last_column, 2))
    cell = sheet.cell(
        row=5, column=1,
        value=f"Department: {context['department']}    |    {revision_label}    |    Date of issue: {context['issue_date']}",
    )
    cell.font = Font(size=9)
    cell.alignment = left

    header_row = 7
    sheet.cell(row=header_row, column=1, value="Sr. No.")
    sheet.cell(row=header_row, column=2, value="Name of the Supervisor")
    sheet.cell(row=header_row + 1, column=1, value="")
    sheet.cell(row=header_row + 1, column=2, value="Session code")
    for index, day in enumerate(dates):
        column = 3 + index
        sheet.cell(row=header_row, column=column, value=day)
        sheet.cell(
            row=header_row + 1, column=column,
            value="/".join(context["codes_by_date"].get(day, [])),
        )
    for row in (header_row, header_row + 1):
        for column in range(1, last_column + 1):
            cell = sheet.cell(row=row, column=column)
            cell.font = Font(bold=True, size=9)
            cell.fill = head_fill
            cell.alignment = center
            cell.border = border

    row_number = header_row + 2
    for roster_row in context["rows"]:
        sheet.cell(row=row_number, column=1, value=roster_row["serial"])
        sheet.cell(row=row_number, column=2, value=roster_row["name"])
        for index, day in enumerate(dates):
            sheet.cell(
                row=row_number, column=3 + index,
                value=roster_row["cells"].get(day, ""),
            )
        for column in range(1, last_column + 1):
            cell = sheet.cell(row=row_number, column=column)
            cell.border = border
            cell.alignment = center if column != 2 else left
            cell.font = Font(size=9)
        row_number += 1

    signature_row = row_number + 3
    labels = ["Senior Supervisor", "Controller of Examinations (COE)", "Principal"]
    for index, label in enumerate(labels):
        column = 1 + index * max(1, last_column // 3)
        column = min(column, last_column)
        sheet.cell(row=signature_row, column=column, value="____________________________")
        sheet.cell(row=signature_row + 1, column=column, value=label)
        sheet.cell(row=signature_row, column=column).alignment = center
        sheet.cell(row=signature_row + 1, column=column).alignment = center

    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 32
    for index in range(len(dates)):
        sheet.column_dimensions[get_column_letter(3 + index)].width = 12
    sheet.freeze_panes = sheet.cell(row=header_row + 2, column=3)
    sheet.print_title_rows = f"1:{header_row + 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.center.text = "Page &P of &N"
    sheet.oddFooter.center.size = 8

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def roster_filename(context, extension):
    role = "junior_supervisor" if context["role"] == "Junior" else "internal_squad"
    revision = context["revision"]
    suffix = f"_revision_{revision}" if revision > 1 else ""
    return f"PICT_{role}_duty_schedule{suffix}.{extension}"
