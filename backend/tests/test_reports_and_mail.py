"""Tests for signed Word/Excel reports and email delivery formatting."""

import os
import sys
import tempfile
import unittest
from datetime import datetime
from email import policy
from email.parser import BytesParser
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from docx import Document
from openpyxl import load_workbook

from mailer import (
    _open_smtp,
    build_daily_change_message,
    build_faculty_message,
    mail_config,
    send_messages,
)
from reports import (
    daily_change_context,
    faculty_report_context,
    render_daily_change_docx,
    render_daily_change_xlsx,
    render_faculty_docx,
    render_faculty_xlsx,
)

REPORT_ENV = {
    "INSTITUTE_NAME": "Pune Institute of Computer Technology (PICT)",
    "COE_NAME": "Dr. COE Test",
    "COE_TITLE": "Controller of Examinations (COE)",
    "PRINCIPAL_NAME": "Dr. Principal Test",
    "PRINCIPAL_TITLE": "Principal",
    "COE_SIGNATURE_IMAGE": "",
    "PRINCIPAL_SIGNATURE_IMAGE": "",
}


def all_docx_text(payload):
    document = Document(BytesIO(payload))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


class SignedReportTests(unittest.TestCase):
    def setUp(self):
        self.timetable = {
            "_id": "tt-test-1",
            "department": "IT",
            "version": 3,
            "revision": 2,
        }
        self.teacher = {
            "teacher_id": "T1",
            "name": "Teacher One",
            "email": "t1@pict.edu",
            "department": "IT",
            "role": "Junior",
        }
        self.duties = [{
            "duty_id": "d1",
            "date": "01-Mar-2026",
            "slot": "1",
            "session": "Morning",
            "course_id": "C1",
            "year": "SY",
            "room": "R1",
            "role": "Junior",
        }]

    def test_faculty_word_and_excel_reports_include_signature_blocks(self):
        with patch.dict(os.environ, REPORT_ENV, clear=False):
            context = faculty_report_context(
                self.timetable, self.teacher, self.duties,
                generated_at=datetime(2026, 3, 1, 10, 30),
            )
            docx_text = all_docx_text(render_faculty_docx(context))
            self.assertIn("Faculty Invigilation Duty Report", docx_text)
            self.assertIn("Teacher One (T1)", docx_text)
            self.assertIn("01-Mar-2026", docx_text)
            self.assertIn("DIGITALLY SIGNED", docx_text)
            self.assertIn("Dr. COE Test", docx_text)
            self.assertIn("Dr. Principal Test", docx_text)

            workbook = load_workbook(BytesIO(render_faculty_xlsx(context)))
            sheet = workbook["Duty Report"]
            values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
            self.assertIn("Faculty Invigilation Duty Report", values)
            self.assertIn("DIGITALLY SIGNED", values)
            self.assertIn("Dr. COE Test", values)
            self.assertIn("Dr. Principal Test", values)

    def test_daily_change_word_and_excel_reports_include_both_faculties_and_signatures(self):
        rows = [
            {
                "slot": "1", "session": "Morning", "course_id": "C1", "year": "SY",
                "room": "R1", "previous_teacher_name": "Teacher One (T1)",
                "new_teacher_name": "Teacher Two (T2)", "change": "Mutual swap",
                "reason": "Medical appointment", "approved_at": "01-Mar-2026, 10:30 AM",
                "revision": 2,
            },
            {
                "slot": "2", "session": "Afternoon", "course_id": "C2", "year": "TY",
                "room": "R2", "previous_teacher_name": "Teacher Two (T2)",
                "new_teacher_name": "Teacher One (T1)", "change": "Mutual swap",
                "reason": "Medical appointment", "approved_at": "01-Mar-2026, 10:30 AM",
                "revision": 2,
            },
        ]
        with patch.dict(os.environ, REPORT_ENV, clear=False):
            context = daily_change_context(
                self.timetable, "01-Mar-2026", rows,
                generated_at=datetime(2026, 3, 1, 10, 35),
            )
            docx_text = all_docx_text(render_daily_change_docx(context))
            self.assertIn("Daily Faculty Duty Change Report", docx_text)
            self.assertIn("Teacher One (T1)", docx_text)
            self.assertIn("Teacher Two (T2)", docx_text)
            self.assertIn("DIGITALLY SIGNED", docx_text)
            self.assertIn("Dr. COE Test", docx_text)

            workbook = load_workbook(BytesIO(render_daily_change_xlsx(context)))
            sheet = workbook["Daily Changes"]
            values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
            self.assertIn("Daily Faculty Duty Change Report — 01-Mar-2026", values)
            self.assertIn("Dr. Principal Test", values)


class MessageFormattingTests(unittest.TestCase):
    def setUp(self):
        self.timetable = {"_id": "tt-test-1", "department": "IT", "version": 3, "revision": 2}
        self.teacher = {
            "teacher_id": "T1", "name": "Teacher One", "email": "t1@pict.edu",
            "department": "IT", "role": "Junior",
        }
        duty = {
            "duty_id": "d1", "date": "01-Mar-2026", "slot": "1", "session": "Morning",
            "course_id": "C1", "year": "SY", "room": "R1", "role": "Junior",
        }
        with patch.dict(os.environ, REPORT_ENV, clear=False):
            self.context = faculty_report_context(
                self.timetable, self.teacher, [duty],
                generated_at=datetime(2026, 3, 1, 10, 30),
            )
            self.docx = render_faculty_docx(self.context)
            self.xlsx = render_faculty_xlsx(self.context)
        self.report = {
            "context": self.context,
            "buffers": {"docx": self.docx, "xlsx": self.xlsx},
            "filenames": {
                "docx": "Faculty_Duty_T1_IT_Rev2.docx",
                "xlsx": "Faculty_Duty_T1_IT_Rev2.xlsx",
            },
        }
        self.mail_env = {
            "MAIL_ENABLED": "true",
            "MAIL_TEST_MODE": "true",
            "MAIL_TEST_RECIPIENT": "dhairyajotwani5@gmail.com",
            "SMTP_HOST": "smtp.example.com",
            "SMTP_PORT": "587",
            "SMTP_USERNAME": "examcell@pict.edu",
            "SMTP_PASSWORD": "test-password",
            "SMTP_FROM_EMAIL": "examcell@pict.edu",
            "SMTP_FROM_NAME": "PICT Examination Cell",
            "SMTP_USE_TLS": "true",
            "SMTP_USE_SSL": "false",
            "MAIL_DRY_RUN": "false",
            "REPORT_RECIPIENTS": "coe@pict.edu,principal@pict.edu",
        }
        self.mail_env.update(REPORT_ENV)

    def test_faculty_message_redirects_to_test_recipient_and_attaches_both_reports(self):
        with patch.dict(os.environ, self.mail_env, clear=False):
            config = mail_config()
            message, meta = build_faculty_message(self.timetable, self.teacher, self.report, config)
        self.assertEqual(message["To"], "dhairyajotwani5@gmail.com")
        self.assertIn("[TEST to t1@pict.edu]", message["Subject"])
        self.assertEqual(meta["intended_to"], "t1@pict.edu")
        self.assertEqual(meta["delivered_to"], "dhairyajotwani5@gmail.com")

        parsed = BytesParser(policy=policy.default).parsebytes(message.as_bytes())
        self.assertEqual(
            [part.get_filename() for part in parsed.iter_attachments()],
            ["Faculty_Duty_T1_IT_Rev2.docx", "Faculty_Duty_T1_IT_Rev2.xlsx"],
        )
        html_body = parsed.get_body(preferencelist=("html",)).get_content()
        self.assertIn("DIGITALLY SIGNED", html_body)
        self.assertIn("Dr. COE Test", html_body)
        self.assertIn("Dr. Principal Test", html_body)

    def test_faculty_message_uses_actual_email_when_demo_mode_is_off(self):
        env = {**self.mail_env, "MAIL_TEST_MODE": "false"}
        with patch.dict(os.environ, env, clear=False):
            config = mail_config()
            message, meta = build_faculty_message(self.timetable, self.teacher, self.report, config)
        self.assertEqual(message["To"], "t1@pict.edu")
        self.assertEqual(meta["intended_to"], "t1@pict.edu")
        self.assertEqual(meta["delivered_to"], "t1@pict.edu")
        self.assertNotIn("[TEST to", message["Subject"])

    def test_daily_change_message_attaches_word_and_excel_for_each_date(self):
        rows = [{
            "slot": "1", "session": "Morning", "course_id": "C1", "year": "SY",
            "room": "R1", "previous_teacher_name": "Teacher One (T1)",
            "new_teacher_name": "Teacher Two (T2)", "change": "Mutual swap",
            "reason": "Medical appointment", "approved_at": "01-Mar-2026, 10:30 AM",
            "revision": 2,
        }]
        with patch.dict(os.environ, REPORT_ENV, clear=False):
            context = daily_change_context(
                self.timetable, "01-Mar-2026", rows,
                generated_at=datetime(2026, 3, 1, 10, 35),
            )
            report = {
                "context": context,
                "buffers": {
                    "docx": render_daily_change_docx(context),
                    "xlsx": render_daily_change_xlsx(context),
                },
                "filenames": {
                    "docx": "Daily_Change_2026-03-01_IT_Rev2.docx",
                    "xlsx": "Daily_Change_2026-03-01_IT_Rev2.xlsx",
                },
            }
        with patch.dict(os.environ, self.mail_env, clear=False):
            config = mail_config()
            message, meta = build_daily_change_message([report], config)
        self.assertEqual(message["To"], "dhairyajotwani5@gmail.com")
        self.assertIn("Daily Faculty Duty Change Report", message["Subject"])
        self.assertEqual(
            meta["attachments"],
            [
                "Daily_Change_2026-03-01_IT_Rev2.docx",
                "Daily_Change_2026-03-01_IT_Rev2.xlsx",
            ],
        )
        parsed = BytesParser(policy=policy.default).parsebytes(message.as_bytes())
        self.assertEqual(len(list(parsed.iter_attachments())), 2)

    def test_dry_run_writes_eml_instead_of_contacting_smtp(self):
        with tempfile.TemporaryDirectory() as output_dir, patch.dict(
            os.environ,
            {
                **self.mail_env,
                "MAIL_DRY_RUN": "true",
                "REPORT_OUTPUT_DIR": output_dir,
            },
            clear=False,
        ):
            config = mail_config()
            message, meta = build_faculty_message(self.timetable, self.teacher, self.report, config)
            results = send_messages([(message, meta)], config)
            self.assertEqual(results[0]["status"], "dry_run")
            eml_path = Path(results[0]["eml_path"])
            self.assertTrue(eml_path.is_file())
            self.assertGreater(eml_path.stat().st_size, 1000)

    def test_send_messages_uses_one_smtp_connection_for_batch(self):
        with patch.dict(
            os.environ,
            {**self.mail_env, "MAIL_DRY_RUN": "false"},
            clear=False,
        ), patch("mailer._open_smtp") as open_smtp:
            server = MagicMock()
            open_smtp.return_value = server
            config = mail_config()
            first, first_meta = build_faculty_message(
                self.timetable, self.teacher, self.report, config
            )
            results = send_messages([(first, first_meta), (first, first_meta)], config)
            self.assertEqual([item["status"] for item in results], ["sent", "sent"])
            self.assertEqual(open_smtp.call_count, 1)
            self.assertEqual(server.send_message.call_count, 2)

    def test_smtp_connection_uses_starttls_and_configured_login(self):
        with patch.dict(
            os.environ,
            {**self.mail_env, "MAIL_DRY_RUN": "false"},
            clear=False,
        ), patch("mailer.smtplib.SMTP") as smtp_class:
            server = MagicMock()
            smtp_class.return_value = server
            config = mail_config()
            opened = _open_smtp(config)
        self.assertIs(opened, server)
        smtp_class.assert_called_once_with("smtp.example.com", 587, timeout=30)
        server.starttls.assert_called_once()
        server.login.assert_called_once_with("examcell@pict.edu", "test-password")
        server.ehlo.assert_called()


if __name__ == "__main__":
    unittest.main()
